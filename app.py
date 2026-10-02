"""
GQ SHOP Mail Report - Streamlit App
=====================================
起動方法:
    streamlit run app.py
"""

import re
import io
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================
# 定数定義
# ============================================================

CLASSIFICATION_RULES = {
    "商品ページ離脱":     ["お気に入りの商品"],
    "チェックアウト離脱": ["お買い忘れ"],
    "カゴ落ち":           ["カートに商品が残っています"],
    "ゴルフセグメント":   ["Cph", "Penguin", "Munsingwear", "ゴルフ", "GOLF", "BRIEFING GOLF"],
    "除外":               ["重要", "会員サービスリニューアル"],
}

REMARK_MAP = {
    "商品ページ離脱":     "商品ページ離脱",
    "チェックアウト離脱": "チェックアウト離脱",
    "カゴ落ち":           "カゴ落ち（カートページ離脱）",
    "ゴルフセグメント":   "ゴルフセグメント",
    "通常キャンペーン":   "",
    "除外":               "除外",
}

NORMAL_CATEGORIES     = ["通常キャンペーン", "ゴルフセグメント"]
AUTOMATION_CATEGORIES = ["商品ページ離脱", "チェックアウト離脱", "カゴ落ち"]

HEADER_KEYWORDS = {
    "キャンペーンアクティビティ", "チャネル", "タイプ", "セッション",
    "売上", "注文数", "コンバージョン率", "費用", "ROAS", "CPA", "CTR",
    "AOV", "新規のお客様からの注文", "リピーターのお客様からの注文",
}

# ============================================================
# 数値変換関数
# ============================================================

def clean_money(v):
    if v is None:
        return 0.0
    s = str(v).strip()
    if s in ("—", "-", ""):
        return 0.0
    cleaned = re.sub(r"[￥¥,\s]", "", s)
    try:
        return float(cleaned) if cleaned else 0.0
    except ValueError:
        return 0.0

def clean_percent(v):
    if v is None:
        return 0.0
    s = str(v).strip()
    if s in ("—", "-", ""):
        return 0.0
    cleaned = re.sub(r"[％%\s]", "", s)
    try:
        return float(cleaned) / 100.0 if cleaned else 0.0
    except ValueError:
        return 0.0

def clean_number(v):
    if v is None:
        return 0.0
    s = str(v).strip()
    if s in ("—", "-", ""):
        return 0.0
    cleaned = re.sub(r"[,\s]", "", s)
    try:
        return float(cleaned) if cleaned else 0.0
    except ValueError:
        return 0.0

# ============================================================
# 列構成の自動検出・パース
# ============================================================

def detect_fields_per_campaign(filtered_lines):
    channel_indices = [i for i, l in enumerate(filtered_lines) if "Shopify Messaging" in l]
    if len(channel_indices) >= 2:
        from collections import Counter
        gaps = [channel_indices[j+1] - channel_indices[j] for j in range(len(channel_indices)-1)]
        return Counter(gaps).most_common(1)[0][0]
    for n in [8, 9, 10, 11, 12, 13, 14]:
        if len(filtered_lines) % n == 0:
            return n
    return 8

def parse_raw_text(raw_text):
    lines = [l.strip() for l in raw_text.strip().splitlines()]
    filtered = [l for l in lines if l and l not in HEADER_KEYWORDS]

    warnings = []
    records = []
    N = detect_fields_per_campaign(filtered)
    total = len(filtered)

    warnings.append(f"ℹ️ 1キャンペーンあたり {N} 列として処理します（全{total}行）")
    if total % N != 0:
        warnings.append(f"⚠️ 行数（{total}行）が{N}の倍数ではありません。不完全なデータがある可能性があります。")

    i = 0
    while i + N <= total:
        chunk = filtered[i:i + N]
        try:
            record = {
                "キャンペーンアクティビティ": chunk[0],
                "チャネル": chunk[1],
                "タイプ":   chunk[2],
            }
            data_chunk = chunk[3:]
            n = len(data_chunk)
            if n == 5:
                keys = ["セッション_raw", "売上_raw", "注文数_raw", "CVR_raw", "CTR_raw"]
            elif n == 6:
                keys = ["セッション_raw", "売上_raw", "注文数_raw", "CVR_raw", "費用_raw", "CTR_raw"]
            elif n == 8:
                keys = ["セッション_raw", "売上_raw", "注文数_raw", "CVR_raw",
                        "費用_raw", "ROAS_raw", "CPA_raw", "CTR_raw"]
            elif n == 11:
                keys = ["セッション_raw", "売上_raw", "注文数_raw", "CVR_raw",
                        "費用_raw", "ROAS_raw", "CPA_raw", "CTR_raw",
                        "AOV_raw", "新規_raw", "リピーター_raw"]
            else:
                base = ["セッション_raw", "売上_raw", "注文数_raw", "CVR_raw",
                        "費用_raw", "ROAS_raw", "CPA_raw", "CTR_raw",
                        "AOV_raw", "新規_raw", "リピーター_raw"]
                keys = base[:n]

            for k, v in zip(keys, data_chunk):
                record[k] = v
            for key in ["セッション_raw", "売上_raw", "注文数_raw", "CVR_raw", "CTR_raw"]:
                if key not in record:
                    record[key] = ""
            records.append(record)
        except Exception as e:
            warnings.append(f"⚠️ 行{i}〜{i+N-1} パース失敗: {e}")
        i += N

    if i < total:
        warnings.append(f"⚠️ 末尾{total - i}行が未処理: {filtered[i:]}")

    return pd.DataFrame(records), warnings

# ============================================================
# 分類・集計関数
# ============================================================

def normalize_name(name):
    return re.sub(r"[【】「」『』\[\]()（）]", "", name)

def classify_campaign(name):
    norm = normalize_name(name)
    for kw in CLASSIFICATION_RULES["除外"]:
        if kw in name or kw in norm:
            return "除外"
    for cat in ["商品ページ離脱", "チェックアウト離脱", "カゴ落ち", "ゴルフセグメント"]:
        for kw in CLASSIFICATION_RULES[cat]:
            if kw in name or kw in norm:
                return cat
    return "通常キャンペーン"

def calc_agg(df, label):
    s = df["セッション"].sum()
    r = df["売上高"].sum()
    o = df["注文数"].sum()
    cvr = (o / s) if s > 0 else 0.0
    return {"集計区分": label, "セッション合計": int(s),
            "売上高合計": int(r), "注文数合計": int(o),
            "コンバージョン率": f"{cvr*100:.2f}%"}

def build_summary(df):
    return pd.DataFrame([
        calc_agg(df[df["分類"].isin(NORMAL_CATEGORIES)],     "通常キャンペーン（＋ゴルフセグメント）"),
        calc_agg(df[df["分類"].isin(AUTOMATION_CATEGORIES)], "自動配信系"),
        calc_agg(df[df["分類"] != "除外"],                   "全体（除外を除く）"),
    ])

def build_dataframe(raw_text):
    df_raw, warnings = parse_raw_text(raw_text)
    if df_raw.empty:
        return None, None, None, None, warnings

    df = df_raw.copy()
    df["セッション"]       = df["セッション_raw"].apply(clean_number)
    df["売上高"]           = df["売上_raw"].apply(clean_money)
    df["注文数"]           = df["注文数_raw"].apply(clean_number)
    df["コンバージョン率"] = df["CVR_raw"].apply(clean_percent)
    df["CTR"]              = df["CTR_raw"].apply(clean_percent)
    df.drop(columns=[c for c in df.columns if c.endswith("_raw")], inplace=True)

    df["分類"] = df["キャンペーンアクティビティ"].apply(classify_campaign)
    df["備考"] = df["分類"].map(REMARK_MAP).fillna("")

    disp = df.copy()
    disp["コンバージョン率"] = disp["コンバージョン率"].apply(lambda x: f"{x*100:.2f}%")
    disp["CTR"]              = disp["CTR"].apply(lambda x: f"{x*100:.2f}%")
    disp["セッション"]       = disp["セッション"].astype(int)
    disp["売上高"]           = disp["売上高"].astype(int)
    disp["注文数"]           = disp["注文数"].astype(int)

    cols = ["分類", "キャンペーンアクティビティ", "チャネル", "タイプ",
            "セッション", "売上高", "注文数", "コンバージョン率", "CTR", "備考"]
    disp = disp[cols].reset_index(drop=True)

    normal_df = disp[disp["分類"].isin(NORMAL_CATEGORIES)].reset_index(drop=True)
    auto_df   = disp[disp["分類"].isin(AUTOMATION_CATEGORIES)].reset_index(drop=True)
    summary   = build_summary(df)

    return disp, normal_df, auto_df, summary, warnings

# ============================================================
# Excel生成
# ============================================================

def build_excel(display_df, normal_df, auto_df, summary_df):
    wb = Workbook()
    wb.remove(wb.active)
    h_fill  = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
    h_font  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    h_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    d_font  = Font(name="Arial", size=10)
    thin    = Side(style="thin")
    border  = Border(left=thin, right=thin, top=thin, bottom=thin)
    c_align = Alignment(horizontal="center", vertical="center")
    fill_map = {
        "ゴルフセグメント":   PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA"),
        "商品ページ離脱":     PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC"),
        "チェックアウト離脱": PatternFill("solid", start_color="FCE4D6", end_color="FCE4D6"),
        "カゴ落ち":           PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7"),
        "除外":               PatternFill("solid", start_color="D9D9D9", end_color="D9D9D9"),
    }
    widths = {
        "分類": 16, "キャンペーンアクティビティ": 45, "チャネル": 18,
        "タイプ": 8, "セッション": 10, "売上高": 12, "注文数": 8,
        "コンバージョン率": 14, "CTR": 10, "備考": 22,
        "集計区分": 30, "セッション合計": 14, "売上高合計": 14, "注文数合計": 10,
    }
    num_cols = ("セッション", "売上高", "注文数", "セッション合計", "売上高合計", "注文数合計")
    pct_cols = ("コンバージョン率", "CTR")

    def write_sheet(ws, df, title, green_header=False):
        ws.title = title
        hf = PatternFill("solid", start_color="375623", end_color="375623") if green_header else h_fill
        for ci, col in enumerate(df.columns, 1):
            c = ws.cell(1, ci, col)
            c.font = h_font; c.fill = hf; c.alignment = h_align; c.border = border
        for ri, (_, row) in enumerate(df.iterrows(), 2):
            rf = fill_map.get(row.get("分類", ""))
            for ci, val in enumerate(row, 1):
                c = ws.cell(ri, ci, val)
                c.font = d_font; c.border = border
                if rf:
                    c.fill = rf
                cn = df.columns[ci - 1]
                if cn in num_cols:
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif cn in pct_cols:
                    c.alignment = c_align
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")
        for ci, cn in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(ci)].width = widths.get(cn, 14)
        ws.row_dimensions[1].height = 30

    write_sheet(wb.create_sheet(), display_df, "整形済みデータ")
    write_sheet(wb.create_sheet(), normal_df,  "通常キャンペーン")
    write_sheet(wb.create_sheet(), auto_df,    "自動配信")
    write_sheet(wb.create_sheet(), summary_df, "集計", green_header=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# ============================================================
# Streamlit UI（トップレベルに直接記述）
# ============================================================

st.set_page_config(page_title="GQ SHOP Mail Report", page_icon="📊", layout="wide")
st.title("📊 GQ SHOP Mail Report")
st.caption("Shopify Messaging キャンペーン実績データ 整形・集計ツール")

st.subheader("① 元データを貼り付けてください")
raw_text = st.text_area(
    label="Shopify Messagingからコピーしたデータをそのまま貼り付けてください",
    height=300,
    placeholder="キャンペーンアクティビティ\nチャネル\nタイプ\n...",
)

run = st.button("🚀 レポート生成", type="primary", disabled=not raw_text.strip())

if not run:
    st.info("データを貼り付けて「レポート生成」ボタンを押してください。")
    st.stop()

with st.spinner("処理中..."):
    display_df, normal_df, auto_df, summary_df, warnings = build_dataframe(raw_text)

for w in warnings:
    if w.startswith("ℹ️"):
        st.info(w)
    else:
        st.warning(w)

if display_df is None or display_df.empty:
    st.error("データを読み込めませんでした。フォーマットを確認してください。")
    st.stop()

st.success(f"✅ {len(display_df)} 件のキャンペーンを処理しました")

st.subheader("② 集計サマリー")
cols = st.columns(3)
labels = ["通常キャンペーン（＋ゴルフセグメント）", "自動配信系", "全体（除外を除く）"]
for i, col in enumerate(cols):
    row = summary_df.iloc[i]
    with col:
        st.metric("集計区分",         labels[i])
        st.metric("セッション",       f"{row['セッション合計']:,}")
        st.metric("売上高",           f"¥{row['売上高合計']:,}")
        st.metric("注文数",           f"{row['注文数合計']:,}")
        st.metric("コンバージョン率", row["コンバージョン率"])
st.dataframe(summary_df, use_container_width=True, hide_index=True)

st.subheader("③ 整形済みデータ（全件）")
st.dataframe(display_df, use_container_width=True, hide_index=True)

st.subheader("④ 通常キャンペーン")
if not normal_df.empty:
    st.dataframe(normal_df, use_container_width=True, hide_index=True)
else:
    st.info("該当データなし")

st.subheader("⑤ 自動配信系キャンペーン")
if not auto_df.empty:
    st.dataframe(auto_df, use_container_width=True, hide_index=True)
else:
    st.info("該当データなし")

st.subheader("⑥ ダウンロード")
dl_cols = st.columns(5)

def to_csv(df):
    return df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")

with dl_cols[0]:
    st.download_button("📄 整形済みデータ CSV", to_csv(display_df), "cleaned_campaign_data.csv", "text/csv")
with dl_cols[1]:
    st.download_button("📄 通常キャンペーン CSV", to_csv(normal_df), "normal_campaigns.csv", "text/csv")
with dl_cols[2]:
    st.download_button("📄 自動配信 CSV", to_csv(auto_df), "automation_campaigns.csv", "text/csv")
with dl_cols[3]:
    st.download_button("📄 集計 CSV", to_csv(summary_df), "summary.csv", "text/csv")
with dl_cols[4]:
    excel_bytes = build_excel(display_df, normal_df, auto_df, summary_df)
    st.download_button("📊 Excelレポート (.xlsx)", excel_bytes, "campaign_report.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
